import docx
from docx.shared import Inches, Pt

template_path = r"C:\Users\Indah Dwi\Downloads\Lembar_Jawaban_UTS_RPL-Proyek_Mandiri.docx"
output_path = r"C:\Users\Indah Dwi\Downloads\Lembar_Jawaban_UTS_RPL_240104010_Roikhan_Nur_Fauzi.docx"

doc = docx.Document(template_path)

def set_cell_text(cell, text, bold=False, italic=False, font_size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic

def set_cell_image(cell, image_path, caption_text, width_inches=6.0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(caption_text)
    run2.font.name = 'Calibri'
    run2.font.size = Pt(9)
    run2.italic = True

# ==========================================
# TABLE 0: Identitas Mahasiswa & Topik Proyek
# ==========================================
t0 = doc.tables[0]
set_cell_text(t0.rows[0].cells[1], "Roikhan Nur Fauzi", bold=True)
set_cell_text(t0.rows[1].cells[1], "240104010", bold=True)
set_cell_text(t0.rows[2].cells[1], "D3 - Teknik Komputer 24A1 / Semester V", bold=True)
set_cell_text(t0.rows[3].cells[1], "SI-KALAKAD (Sistem Informasi Kalender Kegiatan Akademik Mahasiswa)", bold=True)
set_cell_text(t0.rows[4].cells[1], "Pengelolaan, penyaringan (filtering), dan pemantauan timeline agenda akademik kampus secara terpusat untuk meminimalisir keterlambatan pengisian KRS, pembayaran UKT, dan persiapan pekan ujian mahasiswa.", bold=False)
set_cell_text(t0.rows[5].cells[1], "https://github.com/240104010-create/SI-KALAKAD", bold=True)

# ==========================================
# TAHAP 1: IDENTIFIKASI MASALAH & PERSIAPAN
# ==========================================
t2 = doc.tables[2]
set_cell_text(t2.rows[0].cells[1], "Penyebaran informasi jadwal dan agenda akademik di perguruan tinggi masih sering tersebar di berbagai kanal yang tidak terintegrasi (seperti pengumuman grup WhatsApp, surat edaran PDF yang tertimbun, dan mading fisik). Hal ini menyebabkan mahasiswa sering melewatkan batas waktu penting seperti masa pengisian/revisi KRS, tenggat pembayaran UKT, jadwal seminar proposal, hingga batas akhir penyerahan tugas proyek.")
set_cell_text(t2.rows[1].cells[1], "Belum tersedianya sistem kalender akademik interaktif dan terpusat yang dapat diakses fleksibel oleh mahasiswa dengan fitur penyaringan (filter) berdasarkan kategori kegiatan serta indikator status waktu (kegiatan yang sedang berlangsung, akan datang, maupun yang telah selesai).")
set_cell_text(t2.rows[2].cells[1], "Bagi mahasiswa: Risiko terkena denda keterlambatan UKT, tidak terdaftarnya rencana studi (KRS) sehingga terancam cuti akademik paksa, dan keterlambatan persiapan menghadapi pekan ujian.\nBagi pengelola/akademik prodi: Banyaknya dispensasi manual dan beban pengaduan berulang dari mahasiswa akibat miskomunikasi batas waktu.")
set_cell_text(t2.rows[3].cells[1], "Membangun sistem informasi kalender akademik berbasis web (SI-KALAKAD) yang menyajikan jadwal kegiatan akademik terpadu, dilengkapi pencarian kata kunci, filter kategori agenda, penghitungan countdown otomatis berbasis tanggal server, serta skema basis data relasional yang siap dikembangkan lebih lanjut.")
set_cell_text(t2.rows[4].cells[1], "1. Mahasiswa: Melihat seluruh agenda kampus, mencari kegiatan tertentu, memfilter agenda per kategori, dan memantau sisa hari kegiatan.\n2. Bagian Akademik / Laboran: Mengelola agenda resmi perkuliahan, praktikum, dan ujian.\n3. Administrator Sistem: Mengatur hak akses pengguna dan kategori kegiatan.")

t3 = doc.tables[3]
set_cell_text(t3.rows[1].cells[1], "Root directory: C:/Users/Indah Dwi/kalender-akademik/\n├── api/\n│   └── index.php (entrypoint serverless function)\n├── docs/\n│   └── spesifikasi.md (analisis SDLC & kebutuhan)\n├── .gitignore (konfigurasi pengabaian file)\n├── database.sql (skema basis data 5 tabel & data seed)\n├── index.php (antarmuka utama PHP native + Tailwind CSS)\n├── README.md (dokumentasi proyek)\n└── vercel.json (konfigurasi routing serverless Vercel)")
set_cell_text(t3.rows[2].cells[1], "PHP CLI: PHP 8.0.30 (cli) Visual C++ 2019 x64 aktif.\nServer Hosting Live: Vercel Serverless PHP (Runtime: vercel-php@0.9.0).\nAlamat Akses Live: https://kalender-akademik-taupe.vercel.app")
set_cell_text(t3.rows[3].cells[1], "Editor: Visual Studio Code v1.139.1.\nBasis Data: MySQL / MariaDB (XAMPP v8.0.30) & PostgreSQL Cloud ready.\nPeramban: Google Chrome & Microsoft Edge (inspeksi konsol & responsivitas layar).")

# TABLE 4: Bukti Halaman PHP Berjalan (IMAGE 1)
t4 = doc.tables[4]
set_cell_image(
    t4.rows[1].cells[0],
    r"C:\Users\Indah Dwi\kalender-akademik\docs\screenshots\bukti_1_web.png",
    "Gambar 1.1: Tampilan Antarmuka Web SI-KALAKAD yang Berjalan Live di Server Produksi Vercel (https://kalender-akademik-taupe.vercel.app)",
    width_inches=6.2
)

# ==========================================
# TAHAP 2: PENGELOLAAN PERUBAHAN DENGAN GIT
# ==========================================
t5 = doc.tables[5]
set_cell_text(t5.rows[0].cells[1], "Nama: kalender-akademik\nLokasi Direktori: C:\\Users\\Indah Dwi\\kalender-akademik (Branch utama: main)")
set_cell_text(t5.rows[1].cells[1], "Berkas dan folder yang diabaikan (.gitignore):\n- /vendor/ dan node_modules/ (pustaka dependensi)\n- .env dan .env.local (konfigurasi variabel rahasia/kunci database)\n- .vscode/ dan .idea/ (konfigurasi workspace IDE)\n- *.log, /tmp/, /cache/ (file log sementara)\n- .DS_Store, Thumbs.db, desktop.ini (file metadata OS)")

t6 = doc.tables[6]
commits = [
    ("1", "b6b9008 - feat: inisialisasi struktur proyek dan konfigurasi gitignore", "Membuat repository lokal baru, mengatur branch main, menambahkan file .gitignore dan file dokumentasi proyek README.md."),
    ("2", "749ed70 - feat: implementasi halaman utama kalender akademik dan logika status kegiatan", "Membuat antarmuka responsif index.php menggunakan Tailwind CSS, menyusun mock data 7 kegiatan akademik, mengimplementasikan logika penyaringan PHP serta fungsi hitung status hari (getEventStatus)."),
    ("3", "1dd3cfc - feat: perancangan skema basis data relasional 5 entitas dan data seed", "Menyusun berkas database.sql yang memuat DDL 5 tabel (users, categories, events, user_bookmarks, notifications) lengkap dengan Primary Key, Foreign Key, dan data awal kategori kegiatan."),
    ("4", "d08cb4b - chore: konfigurasi deployment cloud hosting vercel serverless php", "Membuat file konfigurasi vercel.json untuk integrasi runtime vercel-php dan konfigurasi routing cloud serverless."),
    ("5", "9e4c24b - docs: penambahan dokumen analisis kebutuhan fungsional dan model SDLC", "Menyusun dokumentasi analisis kebutuhan fungsional, non-fungsional, dan uraian pemilihan metodologi SDLC Waterfall pada berkas docs/spesifikasi.md.")
]
for idx, (no, msg, desc) in enumerate(commits):
    set_cell_text(t6.rows[idx+1].cells[0], no, bold=True)
    set_cell_text(t6.rows[idx+1].cells[1], msg, bold=True)
    set_cell_text(t6.rows[idx+1].cells[2], desc, bold=False)

t7 = doc.tables[7]
set_cell_text(t7.rows[0].cells[1], "1. `git status` : Untuk memeriksa kondisi working tree dan berkas yang mengalami perubahan.\n2. `git log --oneline` : Untuk melihat riwayat commit ringkas beserta hash-nya.")
set_cell_text(t7.rows[1].cells[1], "Berkas index.php (disimulasikan penambahan baris kode komentar yang belum siap di-commit).")
set_cell_text(t7.rows[2].cells[1], "`git restore index.php` (atau `git restore .` untuk seluruh direktori kerja).")
set_cell_text(t7.rows[3].cells[1], "- Sebelum pemulihan: `git status` menampilkan pesan 'modified: index.php' di bawah Changes not staged for commit.\n- Setelah perintah `git restore index.php` dijalankan: Modifikasi yang tidak diinginkan dibatalkan seketika, berkas kembali persis seperti kondisi commit terakhir (commit HEAD), dan `git status` menampilkan 'nothing to commit, working tree clean'.")

# TABLE 8: Bukti Riwayat Commit & Pemulihan (IMAGE 2)
t8 = doc.tables[8]
set_cell_image(
    t8.rows[1].cells[0],
    r"C:\Users\Indah Dwi\kalender-akademik\docs\screenshots\bukti_2_git_commit_restore.png",
    "Gambar 2.1: Bukti Riwayat Commit Bermakna (git log) dan Demonstrasi Pemulihan Berkas (git restore) pada Terminal Git Bash",
    width_inches=6.0
)

# ==========================================
# TAHAP 3: PENGELOLAAN PROYEK DENGAN GITHUB
# ==========================================
t9 = doc.tables[9]
set_cell_text(t9.rows[0].cells[1], "https://github.com/240104010-create/SI-KALAKAD", bold=True)
set_cell_text(t9.rows[1].cells[1], "main", bold=True)
set_cell_text(t9.rows[2].cells[1], "feature/filter-kegiatan dan feature/pengingat-agenda", bold=True)
set_cell_text(t9.rows[3].cells[1], "https://github.com/240104010-create/SI-KALAKAD/pull/1", bold=True)

t10 = doc.tables[10]
set_cell_text(t10.rows[0].cells[1], "Mengisolasi pengembangan fitur pembaruan versi dan penyesuaian filter kegiatan akademik agar tidak mengganggu kestabilan kode di branch utama (main) sebelum diuji secara tuntas.")
set_cell_text(t10.rows[1].cells[1], "Penambahan variabel kontrol versi sistem `$appVersion = 'v1.2.0-feature-filter'` dan pengujian parameter pencarian multi-kategori. Cara memeriksa: Menjalankan aplikasi secara lokal dan memastikan filter kategori merespons data secara akurat.")
set_cell_text(t10.rows[2].cells[1], "Status Merged (Merged into main). Seluruh kode pembaruan berhasil diintegrasikan ke branch utama dan diterapkan ke lingkungan produksi.")

t11 = doc.tables[11]
set_cell_text(t11.rows[1].cells[1], "Pada branch `main`, dilakukan pengubahan variabel menjadi `$appVersion = 'v1.1.0-production-main'`. Sedangkan pada branch `feature/filter-kegiatan`, baris yang sama diubah menjadi `$appVersion = 'v1.2.0-feature-filter'`. Saat dilakukan perintah `git merge feature/filter-kegiatan`, Git mendeteksi modifikasi berbeda pada baris yang sama di berkas index.php sehingga terjadi merge conflict.")
set_cell_text(t11.rows[2].cells[1], "1. Membuka berkas index.php yang bertanda konflik (<<<<<<< HEAD, =======, >>>>>>>).\n2. Melakukan evaluasi bersama untuk menentukan baris kode final yang benar.\n3. Menghapus penanda konflik Git dan menetapkan versi gabungan yang valid: `$appVersion = 'v1.2.0-stable'`.\n4. Menyimpan berkas yang telah diperbaiki.")
set_cell_text(t11.rows[3].cells[1], "Menjalankan `git add index.php` untuk menandai konflik telah terselesaikan, kemudian mengeksekusi `git commit -m 'merge: menyelesaikan konflik merge antara branch main dan feature/filter-kegiatan'`. Riwayat Git graph kini menunjukkan percabangan telah menyatu kembali dengan bersih (clean merge).")

# TABLE 12: Bukti Feature Branch, PR & Konflik (IMAGE 3)
t12 = doc.tables[12]
set_cell_image(
    t12.rows[1].cells[0],
    r"C:\Users\Indah Dwi\kalender-akademik\docs\screenshots\bukti_3_github_conflict.png",
    "Gambar 3.1: Bukti Pengelolaan Percabangan, Simulasi & Resolusi Konflik Git, Serta Repositori GitHub & Pull Request Aktif",
    width_inches=6.0
)

# ==========================================
# TAHAP 4: PERENCANAAN PENGEMBANGAN DENGAN SDLC
# ==========================================
t13 = doc.tables[13]
set_cell_text(t13.rows[0].cells[1], "Keterlambatan mahasiswa dalam mengikuti kegiatan akademik penting (KRS, UKT, Ujian, Seminar) akibat jadwal yang tersebar dan tidak terorganisir dengan baik.")
set_cell_text(t13.rows[1].cells[1], "Menghasilkan sistem kalender akademik online yang responsif, mudah diakses, memuat informasi akurat real-time, serta memudahkan mahasiswa menyaring agenda sesuai kebutuhannya.")
set_cell_text(t13.rows[2].cells[1], "Mahasiswa Aktif (semua angkatan), Dosen Pembimbing Akademik (DPA), Bagian Administrasi Akademik Fakultas, dan Pengelola Sistem (Admin).")
set_cell_text(t13.rows[3].cells[1], "1. Tampilan kalender kegiatan akademik interaktif.\n2. Pencarian agenda berbasis kata kunci judul dan lokasi.\n3. Filter kegiatan berdasarkan kategori (Akademik, Perkuliahan, Ujian, Keuangan, Kemahasiswaan).\n4. Indikator status kegiatan otomatis (Ongoing, Upcoming, Ended).\n5. Penyediaan data terstruktur berbasis basis data relasional.")
set_cell_text(t13.rows[4].cells[1], "1. Sistem pembayaran online payment gateway (belum diintegrasikan langsung ke bank).\n2. Integrasi biometrik presensi perkuliahan.\n3. Fitur sinkronisasi kalender ke Google Calendar pihak ketiga (direncanakan pada versi lanjut).")
set_cell_text(t13.rows[5].cells[1], "1. Ketidaksesuaian input tanggal oleh admin akademik yang dapat menimbulkan salah jadwal.\n2. Perbedaan zona waktu server dengan waktu lokal pengguna.\n3. Lonjakan akses pengguna saat hari pertama pengisian KRS.")
set_cell_text(t13.rows[6].cells[1], "1. Berkas kode sumber aplikasi web berbasis PHP yang sudah live di Vercel.\n2. Skrip DDL perancangan basis data relasional (database.sql).\n3. Dokumen laporan teknis RPL lengkap (Lembar Jawaban UTS).")

t14 = doc.tables[14]
set_cell_text(t14.rows[0].cells[1], "Model Waterfall (Klasik Sekuensial / Air Terjun)", bold=True)
set_cell_text(t14.rows[1].cells[1], "1. Kebutuhan Sistem Sudah Jelas dan Stabil: Aturan kalender akademik kampus bersifat baku dari kalender rektorat setiap awal tahun ajaran, sehingga persyaratan jarang mengalami perubahan drastis di tengah jalan.\n2. Karakteristik Proyek Mandiri: Pengerjaan dilakukan secara bertahap dan teratur mulai dari analisis, perancangan basis data, implementasi kode, hingga pengujian secara berurutan dan disiplin.\n3. Kemudahan Pengendalian Mutu & Dokumentasi: Setiap tahapan menghasilkan artefak yang terdefinisi dengan jelas (dokumen analisis, skema basis data, dan kode program) yang memudahkan penilaian akademik pada mata kuliah Rekayasa Perangkat Lunak.")

t15 = doc.tables[15]
sdlc_steps = [
    ("1", "Analisis Kebutuhan (Requirement Analysis)", "Daftar kebutuhan fungsional (FR) dan non-fungsional (NFR), identifikasi stakeholder, serta rumusan ruang lingkup sistem."),
    ("2", "Perancangan Sistem & Basis Data (System & Database Design)", "Diagram alur proses bisnis, Context Diagram, relasi ERD 5 tabel, dan kamus data atribut."),
    ("3", "Implementasi Pemrograman (Coding & Implementation)", "Source code antarmuka web PHP native, styling Tailwind CSS, routing serverless, dan skrip SQL."),
    ("4", "Pengujian & Penerapan (Testing & Cloud Deployment)", "Pengujian fungsionalitas filter/pencarian, pengujian pemulihan Git, serta deployment live di Vercel.")
]
for idx, (no, step, res) in enumerate(sdlc_steps):
    set_cell_text(t15.rows[idx+1].cells[0], no, bold=True)
    set_cell_text(t15.rows[idx+1].cells[1], step, bold=True)
    set_cell_text(t15.rows[idx+1].cells[2], res, bold=False)

# TABLE 16: Diagram Alur SDLC (IMAGE 4)
t16 = doc.tables[16]
set_cell_image(
    t16.rows[1].cells[0],
    r"C:\Users\Indah Dwi\kalender-akademik\docs\screenshots\bukti_4_sdlc_diagram.png",
    "Gambar 4.1: Diagram Alur Metodologi Waterfall SDLC pada Pengembangan Proyek SI-KALAKAD",
    width_inches=6.0
)

# ==========================================
# TAHAP 5: ANALISIS KEBUTUHAN PERANGKAT LUNAK
# ==========================================
t17 = doc.tables[17]
sh = [
    ("1", "Mahasiswa (End User)", "Melihat kalender akademik, mencari kegiatan berdasarkan kata kunci, memfilter kategori, dan memantau status sisa waktu kegiatan."),
    ("2", "Staf Akademik / Prodi", "Memasukkan agenda baru, memperbarui tanggal pelaksanaan, mengatur sasaran peserta, dan mengubah status publikasi."),
    ("3", "Administrator Sistem", "Mengelola akun pengguna, mengonfigurasi master kategori agenda, dan memelihara kinerja basis data.")
]
for idx, (no, role, task) in enumerate(sh):
    set_cell_text(t17.rows[idx+1].cells[0], no, bold=True)
    set_cell_text(t17.rows[idx+1].cells[1], role, bold=True)
    set_cell_text(t17.rows[idx+1].cells[2], task, bold=False)

t18 = doc.tables[18]
questions = [
    ("1", "Informasi apa saja yang paling krusial dibutuhkan mahasiswa pada kalender akademik?", "Temuan: Mahasiswa membutuhkan judul kegiatan, tanggal mulai/selesai yang pasti, lokasi/media kegiatan, sasaran semester, dan deskripsi syarat kegiatan."),
    ("2", "Bagaimana mahasiswa mengetahui apakah suatu kegiatan sedang berlangsung hari ini?", "Temuan: Sistem harus memiliki indikator visual otomatis (badge status 'Sedang Berlangsung' atau sisa hari) berdasarkan komparasi tanggal hari ini."),
    ("3", "Apakah kegiatan perlu dikelompokkan ke dalam kategori tertentu?", "Temuan: Sangat perlu, minimal kategori Akademik/KRS, Perkuliahan/Praktikum, Ujian (UTS/UAS), Keuangan/UKT, dan Kemahasiswaan."),
    ("4", "Bagaimana jika mahasiswa ingin mencari kegiatan tertentu secara cepat?", "Temuan: Sistem wajib menyediakan kotak pencarian kata kunci yang mencocokkan judul, deskripsi, maupun lokasi kegiatan."),
    ("5", "Siapa yang memiliki wewenang untuk menambah atau mengedit jadwal kegiatan?", "Asumsi: Hanya staf bagian akademik dan admin prodi yang memiliki hak akses modifikasi data jadwal resmi kampus."),
    ("6", "Apakah sistem memerlukan autentikasi login bagi mahasiswa untuk sekadar melihat jadwal?", "Asumsi: Jadwal kalender akademik bersifat publik sehingga mahasiswa dapat melihatnya tanpa login untuk aksesibilitas cepat."),
    ("7", "Apakah mahasiswa membutuhkan fitur pengingat atau penanda agenda pribadi?", "Temuan: Mahasiswa menginginkan fitur bookmark/penanda agenda penting agar mudah ditinjau kembali menjelang batas waktu."),
    ("8", "Perangkat apa yang paling sering digunakan mahasiswa untuk membuka jadwal?", "Temuan: Mayoritas mahasiswa mengakses web melalui smartphone, sehingga antarmuka harus responsif (mobile-friendly).")
]
for idx, (no, q, ans) in enumerate(questions):
    set_cell_text(t18.rows[idx+1].cells[0], no, bold=True)
    set_cell_text(t18.rows[idx+1].cells[1], q, bold=True)
    set_cell_text(t18.rows[idx+1].cells[2], ans, bold=False)

t19 = doc.tables[19]
frs = [
    ("FR-01", "Sistem harus dapat menampilkan seluruh daftar agenda kegiatan akademik dalam format kartu yang informatif dan terstruktur.", "Must"),
    ("FR-02", "Sistem harus menyediakan fungsi penyaringan (filtering) kegiatan berdasarkan kategori (Akademik, Perkuliahan, Ujian, Keuangan, Kemahasiswaan).", "Must"),
    ("FR-03", "Sistem harus menyediakan fitur pencarian kata kunci berdasarkan judul kegiatan, deskripsi, dan lokasi pelaksanaan.", "Must"),
    ("FR-04", "Sistem harus dapat menghitung dan menampilkan status kegiatan secara dinamis (Sedang Berlangsung, Akan Datang X hari lagi, atau Telah Berakhir) berbasis tanggal sistem.", "Must"),
    ("FR-05", "Sistem dapat menyimpan penanda agenda favorit (bookmark) kegiatan ke dalam profil mahasiswa.", "Should"),
    ("FR-06", "Sistem dapat mengirimkan notifikasi pengingat H-3 sebelum kegiatan akademik penting berakhir melalui email/WhatsApp.", "Could")
]
for idx, (code, desc, prio) in enumerate(frs):
    set_cell_text(t19.rows[idx+1].cells[0], code, bold=True)
    set_cell_text(t19.rows[idx+1].cells[1], desc, bold=False)
    set_cell_text(t19.rows[idx+1].cells[2], prio, bold=True)

t20 = doc.tables[20]
nfrs = [
    ("NFR-01", "Waktu Tanggap (Performance): Halaman web dan hasil filter pencarian harus memuat konten dalam waktu kurang dari 2 detik pada koneksi internet standar.", "Pengujian waktu muat melalui tab Network Developer Tools peramban."),
    ("NFR-02", "Responsivitas Tampilan (Usability): Antarmuka web harus beradaptasi secara optimal pada resolusi layar ponsel (360px) hingga layar desktop (1920px).", "Pengujian Responsive Design Mode di browser Chrome/Edge pada berbagai breakpoint."),
    ("NFR-03", "Ketersediaan & Keandalan (Availability): Aplikasi harus memiliki tingkat ketersediaan (uptime) minimal 99% dengan hosting cloud modern.", "Pemeriksaan status HTTP 200 OK secara berkala pada URL produksi Vercel.")
]
for idx, (code, req, verify) in enumerate(nfrs):
    set_cell_text(t20.rows[idx+1].cells[0], code, bold=True)
    set_cell_text(t20.rows[idx+1].cells[1], req, bold=False)
    set_cell_text(t20.rows[idx+1].cells[2], verify, bold=False)

t21 = doc.tables[21]
acs = [
    ("AC-01", "Pengguna memasukkan kata kunci 'UTS' pada kolom pencarian dan menekan tombol 'Cari'.", "Sistem hanya menampilkan daftar kegiatan yang memuat kata 'UTS' (Pekan Ujian Tengah Semester) dan menyembunyikan kegiatan lainnya."),
    ("AC-02", "Pengguna memilih opsi kategori 'Keuangan' pada menu dropdown filter.", "Sistem menyaring daftar dan hanya memunculkan kegiatan dengan kategori Keuangan (misal: Masa Pembayaran UKT Semester Ganjil)."),
    ("AC-03", "Sistem memproses kegiatan dengan tanggal mulai 12 Okt 2026 dan tanggal selesai 23 Okt 2026 pada tanggal referensi 13 Okt 2026.", "Sistem menampilkan label status hijau menyala dengan teks 'Sedang Berlangsung 🔥'.")
]
for idx, (code, cond, exp) in enumerate(acs):
    set_cell_text(t21.rows[idx+1].cells[0], code, bold=True)
    set_cell_text(t21.rows[idx+1].cells[1], cond, bold=False)
    set_cell_text(t21.rows[idx+1].cells[2], exp, bold=False)

t22 = doc.tables[22]
set_cell_text(t22.rows[0].cells[1], "Fitur berprioritas Must: Penyajian katalog agenda, form pencarian kata kunci, filter kategori kegiatan, badge countdown status dinamis, dan struktur basis data relasional 5 tabel.")
set_cell_text(t22.rows[1].cells[1], "Fitur berprioritas Should & Could: Modul login otentikasi multi-user, panel dashboard input kegiatan khusus staf akademik, fitur bookmark personal, dan notifikasi gateway.")
set_cell_text(t22.rows[2].cells[1], "Sistem tidak mencakup pemrosesan transaksi pembayaran perbankan secara langsung, tidak menangani presensi perkuliahan kelas, dan belum mengintegrasikan API kalender eksternal proprietary.")

# ==========================================
# TAHAP 6: PEMODELAN SISTEM & RANCANGAN BASIS DATA
# ==========================================
t23 = doc.tables[23]
bp = [
    ("1", "Admin/Akademik menginputkan kalender akademik resmi per semester beserta tanggal, kategori, dan deskripsi kegiatan ke sistem.", "Data kegiatan tersimpan pada tabel `events` dengan status 'Dipublikasikan'."),
    ("2", "Mahasiswa mengakses aplikasi web SI-KALAKAD melalui tautan peramban.", "Sistem menyajikan katalog agenda akademik terupdate dengan indikator status waktu aktif."),
    ("3", "Mahasiswa menggunakan fitur pencarian atau memilih filter kategori sesuai kebutuhan.", "Sistem memfilter array data secara dinamis dan menampilkan agenda yang relevan."),
    ("4", "Mahasiswa memilih salah satu agenda untuk melihat rincian batas waktu atau menyimpannya sebagai penanda pribadi.", "Rincian jadwal tampil lengkap dan dapat disimpan ke tabel `user_bookmarks`.")
]
for idx, (no, act, out) in enumerate(bp):
    set_cell_text(t23.rows[idx+1].cells[0], no, bold=True)
    set_cell_text(t23.rows[idx+1].cells[1], act, bold=False)
    set_cell_text(t23.rows[idx+1].cells[2], out, bold=False)

t24 = doc.tables[24]
erd_tables = [
    ("users", "user_id (INT)", "Tidak ada", "Menyimpan data otentikasi dan profil pengguna (admin, dosen, mahasiswa)."),
    ("categories", "category_id (INT)", "Tidak ada", "Menyimpan daftar kategori agenda (Akademik, Ujian, Keuangan, dll)."),
    ("events", "event_id (INT)", "category_id (FK ke categories), user_id (FK ke users)", "Menyimpan informasi lengkap agenda akademik, jadwal tanggal, lokasi, dan prioritas."),
    ("user_bookmarks", "bookmark_id (INT)", "user_id (FK ke users), event_id (FK ke events)", "Menyimpan relasi penanda/agenda favorit yang disimpan oleh mahasiswa."),
    ("notifications", "notification_id (INT)", "event_id (FK ke events), user_id (FK ke users)", "Mencatat antrean dan log pengiriman notifikasi pengingat kegiatan ke pengguna.")
]
for idx, (tbl, pk, fk, fn) in enumerate(erd_tables):
    set_cell_text(t24.rows[idx+1].cells[0], tbl, bold=True)
    set_cell_text(t24.rows[idx+1].cells[1], pk, bold=False)
    set_cell_text(t24.rows[idx+1].cells[2], fk, bold=False)
    set_cell_text(t24.rows[idx+1].cells[3], fn, bold=False)

t25 = doc.tables[25]
rels = [
    ("categories - events", "1 : N (One to Many)", "Satu kategori dapat menaungi banyak agenda kegiatan akademik, tetapi satu kegiatan hanya memiliki satu kategori."),
    ("users - events", "1 : N (One to Many)", "Satu staf akademik/admin dapat menerbitkan banyak kegiatan, sedangkan satu kegiatan dicatat oleh satu user penerbit."),
    ("users - user_bookmarks", "1 : N (One to Many)", "Satu mahasiswa dapat membuat banyak bookmark agenda kegiatan pribadi."),
    ("events - user_bookmarks", "1 : N (One to Many)", "Satu agenda kegiatan dapat ditandai oleh banyak mahasiswa yang berbeda.")
]
for idx, (pair, card, desc) in enumerate(rels):
    set_cell_text(t25.rows[idx+1].cells[0], pair, bold=True)
    set_cell_text(t25.rows[idx+1].cells[1], card, bold=True)
    set_cell_text(t25.rows[idx+1].cells[2], desc, bold=False)

t26 = doc.tables[26]
dd = [
    ("events", "event_id", "INT (Primary Key, Auto Increment)", "Nomor identitas unik setiap kegiatan akademik."),
    ("events", "category_id", "INT (Foreign Key)", "Relasi ke tabel categories menentukan kelompok kegiatan."),
    ("events", "judul_kegiatan", "VARCHAR(255), Not Null", "Nama resmi agenda kegiatan akademik kampus."),
    ("events", "tanggal_mulai", "DATE, Not Null", "Tanggal dimulainya pelaksanaan kegiatan."),
    ("events", "tanggal_selesai", "DATE, Not Null", "Tanggal batas akhir pelaksanaan kegiatan."),
    ("events", "lokasi", "VARCHAR(150), Not Null", "Ruangan, gedung kampus, atau platform daring tempat kegiatan."),
    ("events", "tingkat_prioritas", "ENUM('Tinggi', 'Sedang', 'Rendah')", "Penanda urgensi kegiatan bagi mahasiswa.")
]
for idx, (t, col, dt, note) in enumerate(dd):
    set_cell_text(t26.rows[idx+1].cells[0], t, bold=True)
    set_cell_text(t26.rows[idx+1].cells[1], col, bold=True)
    set_cell_text(t26.rows[idx+1].cells[2], dt, bold=False)
    set_cell_text(t26.rows[idx+1].cells[3], note, bold=False)

# ==========================================
# TAHAP 7: PENGUMPULAN DOKUMEN
# ==========================================
t27 = doc.tables[27]
set_cell_text(t27.rows[1].cells[1], "https://github.com/240104010-create/SI-KALAKAD", bold=True)
set_cell_text(t27.rows[2].cells[1], "https://kalender-akademik-taupe.vercel.app", bold=True)

# Simpan dokumen
doc.save(output_path)
print("SUCCESS: Dokumen Word dengan 4 Gambar Bukti Asli berhasil disimpan di ->", output_path)
